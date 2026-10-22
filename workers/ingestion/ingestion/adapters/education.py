from __future__ import annotations

import csv
import io
import os
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, ClassVar
from urllib.parse import urlparse

import httpx
from lente_api.models import (
    Dataset,
    EducationEnrollment,
    EducationMetric,
    PublicationRelease,
    School,
    SchoolVersion,
    SourceSnapshot,
)
from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.orm import Session
from tenacity import retry, retry_if_exception_type, stop_after_attempt, wait_exponential

from ingestion.adapters.base import (
    USER_AGENT,
    FetchedPayload,
    SourceAdapter,
    SourceValidationError,
)

CENSO_2024_URL = "https://download.inep.gov.br/dados_abertos/microdados_censo_escolar_2024.zip"
IDEB_2025_URL = (
    "https://download.inep.gov.br/ideb/resultados/divulgacao_anos_iniciais_municipios_2025.zip"
)
ALLOWED_HOST = "download.inep.gov.br"
MISSING = {"", "-", "...", "NA", "N/A", "NULL"}


@dataclass(slots=True)
class SchoolRecord:
    inep_code: str
    municipality_code: str
    year: int
    name: str
    state: str
    dependency_code: str
    operating_status: str
    location_type: str | None
    infrastructure: dict[str, bool | None]
    enrollments: dict[str, int | None]


@dataclass(slots=True)
class MetricRecord:
    municipality_code: str
    year: int
    metric_code: str
    network_code: str
    stage_code: str
    value: Decimal | None
    availability_status: str


def _safe_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip().upper() in MISSING:
        return None
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation as exc:
        raise SourceValidationError(
            f"Valor numérico inválido no Ideb: {value!r}", code="ideb_invalid_decimal"
        ) from exc


def _safe_int(value: Any) -> int | None:
    parsed = _safe_decimal(value)
    return None if parsed is None else int(parsed)


def _binary_flag(value: Any) -> bool | None:
    raw = str(value or "").strip()
    if raw in {"0", "1"}:
        return raw == "1"
    return None


class _InepDownloadAdapter(SourceAdapter):
    source_url: str
    max_download_bytes = 750 * 1024 * 1024

    def __init__(self, *, raw_data_dir: Path, selected_municipalities: set[str]):
        super().__init__(raw_data_dir=raw_data_dir)
        self.selected_municipalities = selected_municipalities

    @retry(
        retry=retry_if_exception_type((httpx.TimeoutException, httpx.TransportError)),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        reraise=True,
    )
    def discover(self) -> dict[str, Any]:
        parsed = urlparse(self.source_url)
        if parsed.scheme != "https" or parsed.hostname != ALLOWED_HOST:
            raise SourceValidationError(
                "Host do recurso Inep não permitido", code="unsafe_source_url"
            )
        with httpx.Client(
            headers={"User-Agent": USER_AGENT}, follow_redirects=True, timeout=60
        ) as client:
            response = client.head(self.source_url)
            response.raise_for_status()
        if urlparse(str(response.url)).hostname != ALLOWED_HOST:
            raise SourceValidationError(
                "Redirecionamento fora do domínio Inep", code="unsafe_redirect"
            )
        length = int(response.headers.get("content-length", "0"))
        if length > self.max_download_bytes:
            raise SourceValidationError(
                "Arquivo do Inep excede o limite", code="download_too_large"
            )
        return {
            "url": str(response.url),
            "status": response.status_code,
            "content_type": response.headers.get("content-type"),
            "content_length": length,
        }

    @retry(
        retry=retry_if_exception_type((httpx.TimeoutException, httpx.TransportError)),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        reraise=True,
    )
    def fetch(self) -> list[FetchedPayload]:
        stage_dir = self.raw_data_dir / "_stage"
        stage_dir.mkdir(parents=True, exist_ok=True)
        file_handle, temporary_name = tempfile.mkstemp(
            prefix=f"{self.dataset_code}-", suffix=".zip", dir=stage_dir
        )
        os.close(file_handle)
        target = Path(temporary_name)
        fetched_at = datetime.now(UTC)
        headers: dict[str, str] = {}
        total = 0
        try:
            with (
                httpx.Client(
                    headers={"User-Agent": USER_AGENT},
                    follow_redirects=True,
                    timeout=httpx.Timeout(180),
                ) as client,
                client.stream("GET", self.source_url) as response,
            ):
                response.raise_for_status()
                if urlparse(str(response.url)).hostname != ALLOWED_HOST:
                    raise SourceValidationError(
                        "Redirecionamento fora do domínio Inep",
                        code="unsafe_redirect",
                    )
                headers = self.selected_headers(response)
                with target.open("wb") as stream:
                    for chunk in response.iter_bytes(1024 * 1024):
                        total += len(chunk)
                        if total > self.max_download_bytes:
                            raise SourceValidationError(
                                "Arquivo do Inep excede o limite",
                                code="download_too_large",
                            )
                        stream.write(chunk)
            if target.read_bytes()[:2] != b"PK":
                raise SourceValidationError("Recurso baixado não é ZIP", code="unexpected_format")
            self._validate_zip(target)
            return [
                FetchedPayload(
                    content=None,
                    file_path=target,
                    url=self.source_url,
                    fetched_at=fetched_at,
                    original_format="zip",
                    response_headers=headers,
                    published_at=self.header_datetime_from_headers(headers, "last-modified"),
                )
            ]
        except Exception:
            target.unlink(missing_ok=True)
            raise

    @staticmethod
    def header_datetime_from_headers(headers: dict[str, str], name: str) -> datetime | None:
        from dateutil.parser import parse

        raw = next((value for key, value in headers.items() if key.lower() == name), None)
        return parse(raw).astimezone(UTC) if raw else None

    @staticmethod
    def _validate_zip(path: Path) -> None:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(item.file_size for item in members) > 2 * 1024**3:
                raise SourceValidationError("ZIP excede limites seguros", code="unsafe_archive")
            if any(
                Path(item.filename).is_absolute() or ".." in Path(item.filename).parts
                for item in members
            ):
                raise SourceValidationError("Caminho inseguro no ZIP", code="unsafe_archive")

    def _release(self, db: Session, snapshot: SourceSnapshot, year: int) -> PublicationRelease:
        release_key = f"{year}-{snapshot.sha256[:16]}"
        release = db.scalar(
            select(PublicationRelease).where(
                PublicationRelease.dataset_id == snapshot.dataset_id,
                PublicationRelease.release_key == release_key,
            )
        )
        if release is not None:
            return release
        db.execute(
            update(PublicationRelease)
            .where(
                PublicationRelease.dataset_id == snapshot.dataset_id,
                PublicationRelease.status == "published",
            )
            .values(status="withdrawn")
        )
        release = PublicationRelease(
            dataset_id=snapshot.dataset_id,
            release_key=release_key,
            status="published",
            reference_start=date(year, 1, 1),
            reference_end=date(year, 12, 31),
            published_at=datetime.now(UTC),
            verified_at=datetime.now(UTC),
            provenance={
                "snapshot_sha256": snapshot.sha256,
                "source_url": snapshot.url,
                "selected_municipalities": sorted(self.selected_municipalities),
            },
        )
        db.add(release)
        db.flush()
        dataset = db.get(Dataset, snapshot.dataset_id)
        if dataset:
            dataset.publication_status = "published"
        return release


class InepCensoEscolarAdapter(_InepDownloadAdapter):
    source_name = "Inep — Censo Escolar 2024"
    dataset_code = "inep_censo_escolar"
    source_url = CENSO_2024_URL
    year = 2024
    stages: ClassVar[dict[str, str]] = {
        "all_basic": "QT_MAT_BAS",
        "early_childhood": "QT_MAT_INF",
        "fundamental": "QT_MAT_FUND",
        "fundamental_early": "QT_MAT_FUND_AI",
        "fundamental_late": "QT_MAT_FUND_AF",
        "secondary": "QT_MAT_MED",
        "technical": "QT_MAT_PROF_TEC",
        "youth_adult": "QT_MAT_EJA",
        "special": "QT_MAT_ESP",
    }
    infrastructure_fields: ClassVar[dict[str, str]] = {
        "potable_water": "IN_AGUA_POTAVEL",
        "public_electricity": "IN_ENERGIA_REDE_PUBLICA",
        "public_sewer": "IN_ESGOTO_REDE_PUBLICA",
        "library": "IN_BIBLIOTECA",
        "science_lab": "IN_LABORATORIO_CIENCIAS",
        "computer_lab": "IN_LABORATORIO_INFORMATICA",
        "sports_court": "IN_QUADRA_ESPORTES",
        "internet": "IN_INTERNET",
        "broadband": "IN_BANDA_LARGA",
    }

    def parse(self, payload: FetchedPayload) -> list[SchoolRecord]:
        records: list[SchoolRecord] = []
        with zipfile.ZipFile(payload.file_path or io.BytesIO(payload.read_bytes())) as archive:
            names = [
                name
                for name in archive.namelist()
                if name.lower().endswith(".csv") and "microdados_ed_basica" in name.lower()
            ]
            if len(names) != 1:
                raise SourceValidationError(
                    "CSV principal do Censo não identificado",
                    code="censo_schema_changed",
                )
            with (
                archive.open(names[0]) as binary,
                io.TextIOWrapper(binary, encoding="cp1252", newline="") as text,
            ):
                reader = csv.DictReader(text, delimiter=";")
                required = {
                    "NU_ANO_CENSO",
                    "SG_UF",
                    "CO_MUNICIPIO",
                    "NO_ENTIDADE",
                    "CO_ENTIDADE",
                    "TP_DEPENDENCIA",
                    "TP_SITUACAO_FUNCIONAMENTO",
                    "TP_LOCALIZACAO",
                    *self.stages.values(),
                    *self.infrastructure_fields.values(),
                }
                missing = required - set(reader.fieldnames or [])
                if missing:
                    raise SourceValidationError(
                        f"Schema do Censo alterado: {sorted(missing)}",
                        code="censo_schema_changed",
                    )
                for row in reader:
                    municipality = str(row["CO_MUNICIPIO"]).strip()
                    if municipality not in self.selected_municipalities:
                        continue
                    records.append(
                        SchoolRecord(
                            inep_code=str(row["CO_ENTIDADE"]).strip(),
                            municipality_code=municipality,
                            year=int(row["NU_ANO_CENSO"]),
                            name=str(row["NO_ENTIDADE"]).strip(),
                            state=str(row["SG_UF"]).strip(),
                            dependency_code=str(row["TP_DEPENDENCIA"]).strip(),
                            operating_status=str(row["TP_SITUACAO_FUNCIONAMENTO"]).strip(),
                            location_type={"1": "urban", "2": "rural"}.get(
                                str(row["TP_LOCALIZACAO"]).strip()
                            ),
                            infrastructure={
                                name: _binary_flag(row[field])
                                for name, field in self.infrastructure_fields.items()
                            },
                            enrollments={
                                name: _safe_int(row[field]) for name, field in self.stages.items()
                            },
                        )
                    )
        return records

    def validate(self, records: list[SchoolRecord]) -> None:
        if not records:
            raise SourceValidationError("Censo sem escolas para o recorte", code="empty_selection")
        if any(item.year != self.year or len(item.inep_code) != 8 for item in records):
            raise SourceValidationError(
                "Ano ou código Inep inválido", code="censo_invalid_identity"
            )
        if len({item.inep_code for item in records}) != len(records):
            raise SourceValidationError("Escola duplicada no Censo", code="censo_duplicate_school")

    def upsert(self, db: Session, records: list[SchoolRecord], snapshot: SourceSnapshot) -> None:
        release = self._release(db, snapshot, self.year)
        if db.scalar(
            select(SchoolVersion.id).where(SchoolVersion.release_id == release.id).limit(1)
        ):
            return
        now = datetime.now(UTC)
        existing = {
            item.inep_code: item
            for item in db.scalars(
                select(School).where(School.inep_code.in_([record.inep_code for record in records]))
            )
        }
        for record in records:
            school = existing.get(record.inep_code)
            if school is None:
                school = School(
                    inep_code=record.inep_code,
                    municipality_ibge_code=record.municipality_code,
                    created_at=now,
                    updated_at=now,
                )
                db.add(school)
            else:
                school.municipality_ibge_code = record.municipality_code
                school.updated_at = now
        # The explicit flush guarantees parent rows exist before the bulk fact inserts.
        db.flush()
        for record in records:
            db.add(
                SchoolVersion(
                    id=uuid.uuid4(),
                    school_inep_code=record.inep_code,
                    municipality_ibge_code=record.municipality_code,
                    census_year=record.year,
                    name=record.name,
                    state=record.state,
                    dependency_code=record.dependency_code,
                    operating_status=record.operating_status,
                    location_type=record.location_type,
                    infrastructure=record.infrastructure,
                    snapshot_id=snapshot.id,
                    release_id=release.id,
                )
            )
            for stage, count in record.enrollments.items():
                db.add(
                    EducationEnrollment(
                        id=uuid.uuid4(),
                        school_inep_code=record.inep_code,
                        municipality_ibge_code=record.municipality_code,
                        census_year=record.year,
                        stage_code=stage,
                        dependency_code=record.dependency_code,
                        enrollment_count=count,
                        availability_status="available" if count is not None else "unavailable",
                        snapshot_id=snapshot.id,
                        release_id=release.id,
                    )
                )
        db.flush()


class InepIdebAdapter(_InepDownloadAdapter):
    source_name = "Inep — Ideb municipal 2025"
    dataset_code = "inep_ideb_municipal"
    source_url = IDEB_2025_URL
    year = 2025
    metrics: ClassVar[dict[str, str]] = {
        "ideb": "VL_OBSERVADO_2025",
        "approval_rate": "VL_INDICADOR_REND_2025",
        "saeb_mathematics": "VL_NOTA_MATEMATICA_2025",
        "saeb_portuguese": "VL_NOTA_PORTUGUES_2025",
        "saeb_mean": "VL_NOTA_MEDIA_2025",
    }
    networks: ClassVar[dict[str, str]] = {
        "Estadual": "state",
        "Municipal": "municipal",
        "Pública": "public",
    }

    def parse(self, payload: FetchedPayload) -> list[MetricRecord]:
        records: list[MetricRecord] = []
        with zipfile.ZipFile(payload.file_path or io.BytesIO(payload.read_bytes())) as archive:
            names = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
            if len(names) != 1:
                raise SourceValidationError(
                    "Planilha XLSX do Ideb não identificada", code="ideb_schema_changed"
                )
            with archive.open(names[0]) as workbook_stream:
                workbook = load_workbook(workbook_stream, read_only=True, data_only=True)
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers: list[str] | None = None
                for index, values in enumerate(rows, start=1):
                    if index == 10:
                        headers = [
                            str(value).strip() if value is not None else "" for value in values
                        ]
                        required = {"CO_MUNICIPIO", "REDE", *self.metrics.values()}
                        if not required.issubset(headers):
                            raise SourceValidationError(
                                "Schema do Ideb alterado", code="ideb_schema_changed"
                            )
                        continue
                    if index <= 10 or headers is None:
                        continue
                    row = dict(zip(headers, values, strict=False))
                    code = str(row.get("CO_MUNICIPIO") or "").strip()
                    code = code.removesuffix(".0")
                    if code not in self.selected_municipalities:
                        continue
                    network = self.networks.get(str(row.get("REDE") or "").strip())
                    if network is None:
                        continue
                    for metric_code, field in self.metrics.items():
                        value = _safe_decimal(row.get(field))
                        records.append(
                            MetricRecord(
                                code,
                                self.year,
                                metric_code,
                                network,
                                "fundamental_early",
                                value,
                                "available" if value is not None else "suppressed",
                            )
                        )
                workbook.close()
        return records

    def validate(self, records: list[MetricRecord]) -> None:
        if not records:
            raise SourceValidationError("Ideb sem linhas para o recorte", code="empty_selection")
        if any(
            item.year != self.year or item.value is not None and item.value < 0 for item in records
        ):
            raise SourceValidationError("Valor ou ano inválido no Ideb", code="ideb_invalid_value")

    def upsert(self, db: Session, records: list[MetricRecord], snapshot: SourceSnapshot) -> None:
        release = self._release(db, snapshot, self.year)
        if db.scalar(
            select(EducationMetric.id).where(EducationMetric.release_id == release.id).limit(1)
        ):
            return
        for record in records:
            db.add(
                EducationMetric(
                    id=uuid.uuid4(),
                    municipality_ibge_code=record.municipality_code,
                    reference_year=record.year,
                    metric_code=record.metric_code,
                    network_code=record.network_code,
                    stage_code=record.stage_code,
                    value=record.value,
                    availability_status=record.availability_status,
                    snapshot_id=snapshot.id,
                    release_id=release.id,
                )
            )
        db.flush()


def rollback_education_release(db: Session, dataset_code: str, release_key: str) -> str:
    dataset = db.scalar(select(Dataset).where(Dataset.code == dataset_code))
    if dataset is None:
        raise ValueError("Dataset educacional não encontrado")
    release = db.scalar(
        select(PublicationRelease).where(
            PublicationRelease.dataset_id == dataset.id,
            PublicationRelease.release_key == release_key,
        )
    )
    if release is None or release.status != "published":
        raise ValueError("Release publicada não encontrada")
    release.status = "withdrawn"
    previous = db.scalar(
        select(PublicationRelease)
        .where(
            PublicationRelease.dataset_id == dataset.id,
            PublicationRelease.status == "withdrawn",
            PublicationRelease.id != release.id,
        )
        .order_by(PublicationRelease.published_at.desc())
        .limit(1)
    )
    if previous:
        previous.status = "published"
        dataset.publication_status = "published"
        result = previous.release_key
    else:
        dataset.publication_status = "not_published"
        result = "none"
    db.commit()
    return result
