from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, ClassVar

from lente_api.models import CatalogEntry, Dataset, SourceSnapshot
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ingestion.adapters.base import (
    FetchedPayload,
    ObservationRecord,
    SourceAdapter,
    SourceValidationError,
)


@dataclass(frozen=True, slots=True)
class MunicVariable:
    indicator_code: str
    column: str


@dataclass(frozen=True, slots=True)
class MunicProfile:
    dataset_code: str
    year: int
    url: str
    resource_version: str
    variables: tuple[MunicVariable, ...]


PROFILES: dict[str, MunicProfile] = {
    "digital": MunicProfile(
        "ibge_munic_2024_digital_government",
        2024,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2024/Base_de_Dados/Base_MUNIC_2024_20251107.xlsx",
        "Base_MUNIC_2024_20251107",
        (
            MunicVariable("municipal_tic_structure", "Mtic06"),
            MunicVariable("municipal_data_center", "Mtic08"),
            MunicVariable("municipal_software_internal", "Mtic091"),
            MunicVariable("municipal_software_public", "Mtic092"),
        ),
    ),
    "governance": MunicProfile(
        "ibge_munic_2024_governance",
        2024,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2024/Base_de_Dados/Base_MUNIC_2024_20251107.xlsx",
        "Base_MUNIC_2024_20251107",
        (
            MunicVariable("municipal_internal_control", "Mgov09"),
            MunicVariable("internal_control_online", "Mgov23"),
        ),
    ),
    "culture": MunicProfile(
        "ibge_munic_2021_culture",
        2021,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2021/Base_de_Dados/Base_MUNIC_2021_20240425.xlsx",
        "Base_MUNIC_2021_20240425",
        (
            MunicVariable("municipal_culture_plan", "Mcul10"),
            MunicVariable("municipal_culture_council", "Mcul19"),
            MunicVariable("municipal_culture_fund", "Mcul33"),
            MunicVariable("public_library_exists", "Mcul3901"),
            MunicVariable("museum_exists", "Mcul3902"),
            MunicVariable("theater_exists", "Mcul3903"),
        ),
    ),
    "participation": MunicProfile(
        "ibge_munic_2021_participation",
        2021,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2021/Base_de_Dados/Base_MUNIC_2021_20240425.xlsx",
        "Base_MUNIC_2021_20240425",
        (
            MunicVariable("culture_plan_civic_participation", "Mcul11"),
            MunicVariable("municipal_culture_council", "Mcul19"),
        ),
    ),
    "environment": MunicProfile(
        "ibge_munic_2020_environment",
        2020,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2020/Base_de_Dados/Base_MUNIC_2020.xlsx",
        "Base_MUNIC_2020",
        (
            MunicVariable("municipal_environment_council", "Mmam10"),
            MunicVariable("environment_council_meetings", "Mmam12"),
            MunicVariable("environment_council_members", "Mmam13"),
            MunicVariable("municipal_environment_fund", "Mmam17"),
        ),
    ),
    "risks": MunicProfile(
        "ibge_munic_2020_risk_management",
        2020,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2020/Base_de_Dados/Base_MUNIC_2020.xlsx",
        "Base_MUNIC_2020",
        (
            MunicVariable("flooding_last_four_years", "Mgrd06"),
            MunicVariable("gradual_flood_last_four_years", "Mgrd08"),
            MunicVariable("flash_flood_last_four_years", "Mgrd11"),
            MunicVariable("landslide_last_four_years", "Mgrd14"),
            MunicVariable("municipal_risk_reduction_plan", "Mgrd177"),
        ),
    ),
    "social": MunicProfile(
        "ibge_munic_2023_social_assistance",
        2023,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2023/Base_de_Dados/Base_MUNIC_2023.xlsx",
        "Base_MUNIC_2023",
        (
            MunicVariable("municipal_social_assistance_plan", "MASS11"),
            MunicVariable("social_plan_council_approved", "MASS12"),
            MunicVariable("municipal_social_assistance_council", "MASS16"),
            MunicVariable("municipal_social_assistance_fund", "MASS22"),
        ),
    ),
    "security": MunicProfile(
        "ibge_munic_2023_public_security",
        2023,
        "https://ftp.ibge.gov.br/Perfil_Municipios/2023/Base_de_Dados/Base_MUNIC_2023.xlsx",
        "Base_MUNIC_2023",
        (
            MunicVariable("municipal_public_security_council", "MSEG07"),
            MunicVariable("community_public_security_council", "MSEG13"),
            MunicVariable("municipal_public_security_fund", "MSEG14"),
            MunicVariable("municipal_public_security_plan", "MSEG15"),
            MunicVariable("municipal_guard", "MSEG27"),
        ),
    ),
}


class IbgeMunicAdapter(SourceAdapter):
    source_name = "ibge_munic"
    dataset_code = ""
    max_resource_bytes: ClassVar[int] = 40_000_000

    def __init__(
        self,
        *,
        raw_data_dir: Path,
        municipalities: set[str],
        profile: str,
    ):
        super().__init__(raw_data_dir=raw_data_dir)
        if profile not in PROFILES:
            raise ValueError(f"Perfil MUNIC desconhecido: {profile}")
        if not municipalities:
            raise ValueError("Ao menos um município deve ser selecionado")
        self.profile_name = profile
        self.profile = PROFILES[profile]
        self.dataset_code = self.profile.dataset_code
        self.source_name = f"ibge_munic_{profile}"
        self.municipalities = municipalities

    def discover(self) -> dict[str, Any]:
        return {
            "survey": "Pesquisa de Informações Básicas Municipais — MUNIC",
            "edition": self.profile.year,
            "resource_version": self.profile.resource_version,
            "municipalities": sorted(self.municipalities),
            "variables": [item.column for item in self.profile.variables],
        }

    def fetch(self) -> list[FetchedPayload]:
        response = self.get(self.profile.url, timeout=180)
        if len(response.content) > self.max_resource_bytes:
            raise SourceValidationError(
                "Planilha MUNIC excede o limite homologado", code="resource_size"
            )
        return [
            FetchedPayload(
                content=response.content,
                url=str(response.url),
                fetched_at=datetime.now(UTC),
                original_format="xlsx",
                response_headers=self.selected_headers(response),
                published_at=self.header_datetime(response, "last-modified"),
                persist_raw=True,
                resource_version=self.profile.resource_version,
            )
        ]

    def parse(self, payload: FetchedPayload) -> list[ObservationRecord]:
        try:
            workbook = load_workbook(
                BytesIO(payload.read_bytes()), read_only=True, data_only=True
            )
        except Exception as exc:
            raise SourceValidationError(
                "Planilha MUNIC inválida", code="invalid_workbook"
            ) from exc
        wanted = {item.column.casefold(): item for item in self.profile.variables}
        sheet = None
        headers: list[str] = []
        for candidate in workbook.worksheets[1:]:
            candidate_headers = [
                str(value).strip() if value is not None else ""
                for value in next(candidate.iter_rows(values_only=True))
            ]
            normalized = {value.casefold() for value in candidate_headers}
            if set(wanted).issubset(normalized):
                sheet = candidate
                headers = candidate_headers
                break
        if sheet is None:
            raise SourceValidationError(
                "Variáveis MUNIC não localizadas", code="schema_drift"
            )
        columns = {header.casefold(): index for index, header in enumerate(headers)}
        municipality_column = next(
            (
                index
                for index, header in enumerate(headers)
                if header.casefold().replace(" ", "") in {"codmun", "codmunic"}
            ),
            None,
        )
        if municipality_column is None:
            raise SourceValidationError(
                "Código municipal ausente na MUNIC", code="schema_drift"
            )
        records: list[ObservationRecord] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            municipality = str(row[municipality_column]).split(".", 1)[0]
            if municipality not in self.municipalities:
                continue
            for normalized_column, variable in wanted.items():
                raw = row[columns[normalized_column]]
                value, status = _normalize_value(raw)
                records.append(
                    ObservationRecord(
                        indicator_code=variable.indicator_code,
                        municipality_ibge_code=municipality,
                        period_start=date(self.profile.year, 1, 1),
                        period_end=date(self.profile.year, 12, 31),
                        value=value,
                        scope={
                            "edition": self.profile.year,
                            "variable": variable.column,
                            "response": None if raw is None else str(raw).strip(),
                            "resource_version": self.profile.resource_version,
                        },
                        availability_status=status,
                    )
                )
        return records

    def validate(self, records: list[ObservationRecord]) -> None:
        expected = len(self.municipalities) * len(self.profile.variables)
        if len(records) != expected:
            raise SourceValidationError(
                "Cobertura MUNIC incompleta",
                code="missing_municipality",
                details={"expected": expected, "received": len(records)},
            )
        keys = [
            (record.municipality_ibge_code, record.indicator_code) for record in records
        ]
        if len(keys) != len(set(keys)):
            raise SourceValidationError(
                "MUNIC contém chaves duplicadas", code="duplicate_key"
            )

    def upsert(
        self, db: Session, records: list[ObservationRecord], snapshot: SourceSnapshot
    ) -> None:
        super().upsert(db, records, snapshot)
        dataset = db.get(Dataset, snapshot.dataset_id)
        if dataset is not None:
            dataset.publication_status = "published"
            entries = db.scalars(
                select(CatalogEntry).where(CatalogEntry.dataset_id == dataset.id)
            )
            for entry in entries:
                entry.review_status = "validated_and_published"
                entry.decision_reason = "Recorte MUNIC validado e publicado."
                entry.last_checked_at = snapshot.fetched_at


def _normalize_value(raw: Any) -> tuple[Decimal | None, str]:
    if raw is None:
        return None, "unavailable"
    if isinstance(raw, bool):
        return Decimal(int(raw)), "available"
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)), "available"
    text = str(raw).strip()
    folded = text.casefold()
    if not text or folded in {"-", "não sabe", "na", "não aplicável"}:
        return None, "unavailable"
    if folded.startswith("sim"):
        return Decimal(1), "available"
    if folded.startswith(("não", "nao")):
        return Decimal(0), "available"
    try:
        return Decimal(text.replace(".", "").replace(",", ".")), "available"
    except InvalidOperation:
        return None, "unavailable"
