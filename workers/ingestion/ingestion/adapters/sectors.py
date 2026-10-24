from __future__ import annotations

import csv
import ftplib
import io
import json
import re
import shutil
import unicodedata
import uuid
from calendar import monthrange
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, ClassVar

import py7zr
from openpyxl import load_workbook
from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from ingestion.adapters.base import FetchedPayload, SourceAdapter, SourceValidationError
from lente_api.models import (
    Dataset,
    EconomicActivity,
    EconomyFact,
    InfrastructureFact,
    InfrastructureService,
    LaborFact,
    SourceSnapshot,
)


def normalize_name(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", ascii_value.upper()).strip()


@dataclass(slots=True)
class EconomyRecord:
    municipality_ibge_code: str
    year: int
    metric_code: str
    activity_code: str | None
    activity_label: str | None
    value: Decimal | None
    unit: str
    availability_status: str
    source_url: str


@dataclass(slots=True)
class LaborRecord:
    municipality_ibge_code: str
    period_start: date
    period_end: date
    section_code: str
    admissions: int
    dismissals: int
    balance: int
    source_url: str


@dataclass(slots=True)
class InfrastructureRecord:
    municipality_ibge_code: str
    reference_date: date
    service_code: str
    service_label: str
    value: Decimal | None
    source_url: str


class IbgePibAdapter(SourceAdapter):
    source_name = "ibge_pib"
    dataset_code = "ibge_pib_municipal"
    base_url = "https://apisidra.ibge.gov.br/values/t/5938"
    variables: ClassVar[dict[str, tuple[str, str | None, str | None]]] = {
        "37": ("gdp_current_prices", None, None),
        "498": ("gross_value_added", "TOTAL", "Total"),
        "513": ("gross_value_added", "AGRICULTURE", "Agropecuária"),
        "517": ("gross_value_added", "INDUSTRY", "Indústria"),
        "6575": (
            "gross_value_added",
            "SERVICES",
            "Serviços exceto administração pública",
        ),
        "525": ("gross_value_added", "PUBLIC_ADMINISTRATION", "Administração pública"),
        "543": ("product_taxes_net_subsidies", None, None),
    }

    def __init__(self, *, raw_data_dir: Path, municipalities: set[str], years: list[int]):
        super().__init__(raw_data_dir=raw_data_dir)
        self.municipalities = municipalities
        self.years = sorted(set(years))

    def discover(self) -> dict[str, Any]:
        return {"table": 5938, "variables": sorted(self.variables), "years": self.years}

    def fetch(self) -> list[FetchedPayload]:
        codes = ",".join(sorted(self.municipalities))
        variables = ",".join(self.variables)
        years = ",".join(str(year) for year in self.years)
        url = f"{self.base_url}/n6/{codes}/v/{variables}/p/{years}?formato=json"
        response = self.get(url, timeout=90)
        return [
            FetchedPayload(
                content=response.content,
                url=str(response.url),
                fetched_at=datetime.now(UTC),
                original_format="json",
                response_headers=self.selected_headers(response),
                persist_raw=True,
                resource_version="sidra-5938",
            )
        ]

    def parse(self, payload: FetchedPayload) -> list[EconomyRecord]:
        try:
            rows = json.loads(payload.read_bytes())
        except json.JSONDecodeError as exc:
            raise SourceValidationError("Resposta SIDRA PIB inválida", code="invalid_json") from exc
        required = {"V", "D1C", "D2C", "D3C", "MN"}
        if not isinstance(rows, list) or len(rows) < 2 or not required.issubset(rows[0]):
            raise SourceValidationError("Schema SIDRA PIB alterado", code="schema_drift")
        records: list[EconomyRecord] = []
        for row in rows[1:]:
            municipality = str(row["D1C"])
            variable = str(row["D2C"])
            year = int(row["D3C"])
            if (
                municipality not in self.municipalities
                or variable not in self.variables
                or year not in self.years
            ):
                raise SourceValidationError("Linha SIDRA fora do recorte", code="domain_mismatch")
            raw = str(row.get("V", "")).strip()
            try:
                value = None if raw in {"", "...", "-", "X"} else Decimal(raw) * Decimal(1000)
            except InvalidOperation as exc:
                raise SourceValidationError("Valor PIB inválido", code="invalid_numeric") from exc
            metric, activity, label = self.variables[variable]
            records.append(
                EconomyRecord(
                    municipality,
                    year,
                    metric,
                    activity,
                    label,
                    value,
                    "BRL nominal",
                    "available" if value is not None else "unavailable",
                    payload.url,
                )
            )
        return records

    def validate(self, records: list[EconomyRecord]) -> None:
        if not records:
            raise SourceValidationError("PIB municipal vazio", code="empty_dataset")
        keys = [(r.municipality_ibge_code, r.year, r.metric_code, r.activity_code) for r in records]
        if len(keys) != len(set(keys)):
            raise SourceValidationError("PIB municipal duplicado", code="duplicate_key")
        for municipality in self.municipalities:
            if not any(
                r.municipality_ibge_code == municipality
                and r.metric_code == "gdp_current_prices"
                and r.value is not None
                for r in records
            ):
                raise SourceValidationError(
                    "PIB total ausente para município solicitado",
                    code="missing_required_metric",
                    details={"municipality": municipality},
                )

    def upsert(self, db: Session, records: list[EconomyRecord], snapshot: SourceSnapshot) -> None:
        activities: dict[str, EconomicActivity] = {}
        for record in records:
            if record.activity_code and record.activity_code not in activities:
                activity = db.scalar(
                    select(EconomicActivity).where(
                        EconomicActivity.classification_system == "IBGE_PIB_SECTOR",
                        EconomicActivity.code == record.activity_code,
                    )
                )
                if activity is None:
                    activity = EconomicActivity(
                        id=uuid.uuid4(),
                        classification_system="IBGE_PIB_SECTOR",
                        code=record.activity_code,
                        label=record.activity_label or record.activity_code,
                        level="broad_sector",
                    )
                    db.add(activity)
                    db.flush()
                activities[record.activity_code] = activity
            key = ":".join(
                (
                    record.municipality_ibge_code,
                    str(record.year),
                    record.metric_code,
                    record.activity_code or "NONE",
                )
            )
            previous = db.scalar(
                select(EconomyFact)
                .where(
                    EconomyFact.source_system == "sidra_pib",
                    EconomyFact.source_key == key,
                )
                .order_by(desc(EconomyFact.revision))
                .limit(1)
            )
            if (
                previous
                and previous.value == record.value
                and previous.availability_status == record.availability_status
            ):
                continue
            db.add(
                EconomyFact(
                    id=uuid.uuid4(),
                    source_system="sidra_pib",
                    source_key=key,
                    municipality_ibge_code=record.municipality_ibge_code,
                    reference_year=record.year,
                    metric_code=record.metric_code,
                    activity_id=activities.get(record.activity_code).id
                    if record.activity_code
                    else None,
                    value=record.value,
                    unit=record.unit,
                    price_basis="current_prices",
                    availability_status=record.availability_status,
                    source_document_url=record.source_url,
                    snapshot_id=snapshot.id,
                    revision=1 if previous is None else previous.revision + 1,
                )
            )
        db.flush()
        dataset = db.get(Dataset, snapshot.dataset_id)
        if dataset is not None:
            dataset.publication_status = "published"


class NovoCagedAdapter(SourceAdapter):
    source_name = "novo_caged"
    dataset_code = "novo_caged_movements"
    ftp_host = "ftp.mtps.gov.br"
    section_labels: ClassVar[dict[str, str]] = {
        "A": "Agricultura",
        "B": "Indústrias extrativas",
        "C": "Indústrias de transformação",
        "D": "Eletricidade e gás",
        "E": "Água, esgoto e resíduos",
        "F": "Construção",
        "G": "Comércio",
        "H": "Transporte e armazenagem",
        "I": "Alojamento e alimentação",
        "J": "Informação e comunicação",
        "K": "Atividades financeiras",
        "L": "Atividades imobiliárias",
        "M": "Atividades profissionais",
        "N": "Atividades administrativas",
        "O": "Administração pública",
        "P": "Educação",
        "Q": "Saúde humana e serviços sociais",
        "R": "Artes e recreação",
        "S": "Outras atividades de serviços",
        "T": "Serviços domésticos",
        "U": "Organismos internacionais",
        "Z": "Não informado",
    }

    def __init__(
        self,
        *,
        raw_data_dir: Path,
        municipality_crosswalk: dict[str, str],
        year: int = 2025,
        month: int = 12,
    ):
        super().__init__(raw_data_dir=raw_data_dir)
        self.crosswalk = municipality_crosswalk
        self.year, self.month = year, month
        self.competence = f"{year}{month:02d}"
        self.remote_dir = f"/pdet/microdados/NOVO CAGED/{year}/{self.competence}"
        self.filename = f"CAGEDMOV{self.competence}.7z"

    @property
    def source_url(self) -> str:
        return f"ftp://{self.ftp_host}{self.remote_dir.replace(' ', '%20')}/{self.filename}"

    def discover(self) -> dict[str, Any]:
        with ftplib.FTP(self.ftp_host, timeout=30) as ftp:
            ftp.login()
            ftp.set_pasv(True)
            ftp.cwd(self.remote_dir)
            size = ftp.size(self.filename)
        if not size or size > 500_000_000:
            raise SourceValidationError("Tamanho CAGED fora do limite", code="resource_size")
        return {
            "competence": self.competence,
            "filename": self.filename,
            "compressed_bytes": size,
        }

    def fetch(self) -> list[FetchedPayload]:
        stage = self.raw_data_dir / "_stage"
        stage.mkdir(parents=True, exist_ok=True)
        target = stage / f"{uuid.uuid4()}-{self.filename}"
        try:
            with ftplib.FTP(self.ftp_host, timeout=120) as ftp:
                ftp.login()
                ftp.set_pasv(True)
                ftp.cwd(self.remote_dir)
                with target.open("wb") as stream:
                    ftp.retrbinary(f"RETR {self.filename}", stream.write, blocksize=1024 * 1024)
            return [
                FetchedPayload(
                    content=None,
                    file_path=target,
                    url=self.source_url,
                    fetched_at=datetime.now(UTC),
                    original_format="7z",
                    persist_raw=False,
                    resource_version=self.competence,
                )
            ]
        except Exception:
            target.unlink(missing_ok=True)
            raise

    def parse(self, payload: FetchedPayload) -> list[LaborRecord]:
        if payload.file_path is None:
            raise SourceValidationError("Arquivo CAGED ausente", code="empty_payload")
        extract_dir = payload.file_path.parent / f"extract-{uuid.uuid4()}"
        extract_dir.mkdir()
        try:
            with py7zr.SevenZipFile(payload.file_path, "r") as archive:
                members = archive.list()
                expected = f"CAGEDMOV{self.competence}.txt"
                if (
                    len(members) != 1
                    or Path(members[0].filename).name != expected
                    or members[0].uncompressed > 2_000_000_000
                ):
                    raise SourceValidationError("Conteúdo 7z CAGED inesperado", code="schema_drift")
                archive.extractall(path=extract_dir)
            text_path = extract_dir / expected
            totals: dict[tuple[str, str], list[int]] = defaultdict(lambda: [0, 0, 0])
            with text_path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream, delimiter=";")
                required = {"competênciamov", "município", "seção", "saldomovimentação"}
                if reader.fieldnames is None or not required.issubset(reader.fieldnames):
                    raise SourceValidationError(
                        "Cabeçalho Novo CAGED alterado",
                        code="schema_drift",
                        details={"header": reader.fieldnames},
                    )
                for row in reader:
                    external = str(row["município"]).strip()
                    municipality = self.crosswalk.get(external)
                    if municipality is None:
                        continue
                    if str(row["competênciamov"]).strip() != self.competence:
                        raise SourceValidationError(
                            "Competência CAGED inesperada", code="period_mismatch"
                        )
                    section = str(row["seção"]).strip().upper() or "Z"
                    movement = int(row["saldomovimentação"])
                    if movement not in {-1, 1}:
                        raise SourceValidationError(
                            "Movimentação CAGED inválida", code="invalid_numeric"
                        )
                    bucket = totals[(municipality, section)]
                    if movement == 1:
                        bucket[0] += 1
                    else:
                        bucket[1] += 1
                    bucket[2] += movement
            last_day = date(self.year, self.month, monthrange(self.year, self.month)[1])
            return [
                LaborRecord(
                    municipality,
                    date(self.year, self.month, 1),
                    last_day,
                    section,
                    values[0],
                    values[1],
                    values[2],
                    payload.url,
                )
                for (municipality, section), values in sorted(totals.items())
            ]
        finally:
            shutil.rmtree(extract_dir, ignore_errors=True)
            payload.file_path.unlink(missing_ok=True)

    def validate(self, records: list[LaborRecord]) -> None:
        if not records:
            raise SourceValidationError(
                "Nenhuma movimentação CAGED no recorte", code="empty_dataset"
            )
        if any(r.balance != r.admissions - r.dismissals for r in records):
            raise SourceValidationError("Saldo CAGED não conciliado", code="reconciliation_failed")

    def upsert(self, db: Session, records: list[LaborRecord], snapshot: SourceSnapshot) -> None:
        for record in records:
            activity = db.scalar(
                select(EconomicActivity).where(
                    EconomicActivity.classification_system == "CNAE2_SECTION",
                    EconomicActivity.code == record.section_code,
                )
            )
            if activity is None:
                activity = EconomicActivity(
                    id=uuid.uuid4(),
                    classification_system="CNAE2_SECTION",
                    code=record.section_code,
                    label=self.section_labels.get(record.section_code, "Não informado"),
                    level="section",
                )
                db.add(activity)
                db.flush()
            key = f"{record.municipality_ibge_code}:{self.competence}:{record.section_code}"
            previous = db.scalar(
                select(LaborFact)
                .where(LaborFact.source_system == "novo_caged", LaborFact.source_key == key)
                .order_by(desc(LaborFact.revision))
                .limit(1)
            )
            if previous and (
                previous.admissions,
                previous.dismissals,
                previous.balance,
            ) == (record.admissions, record.dismissals, record.balance):
                continue
            db.add(
                LaborFact(
                    id=uuid.uuid4(),
                    source_system="novo_caged",
                    source_key=key,
                    municipality_ibge_code=record.municipality_ibge_code,
                    period_start=record.period_start,
                    period_end=record.period_end,
                    activity_id=activity.id,
                    admissions=record.admissions,
                    dismissals=record.dismissals,
                    balance=record.balance,
                    availability_status="available",
                    source_document_url=record.source_url,
                    snapshot_id=snapshot.id,
                    revision=1 if previous is None else previous.revision + 1,
                )
            )
        db.flush()
        dataset = db.get(Dataset, snapshot.dataset_id)
        if dataset is not None:
            dataset.publication_status = "published"


class SenatranFleetAdapter(SourceAdapter):
    source_name = "senatran_fleet"
    dataset_code = "senatran_fleet"
    resource_url = "https://www.gov.br/transportes/pt-br/assuntos/transito/conteudo-Senatran/frota-de-veiculos-2026-2/Frota_por_municipio_e_tipo_Setembro_2026.xlsx/@@download/file"

    def __init__(self, *, raw_data_dir: Path, municipality_names: dict[tuple[str, str], str]):
        super().__init__(raw_data_dir=raw_data_dir)
        self.municipality_names = municipality_names

    def discover(self) -> dict[str, Any]:
        return {
            "reference": "2026-09",
            "resource": self.resource_url,
            "municipalities": len(self.municipality_names),
        }

    def fetch(self) -> list[FetchedPayload]:
        response = self.get(self.resource_url, timeout=90)
        if len(response.content) > 20_000_000:
            raise SourceValidationError("Planilha SENATRAN excede limite", code="resource_size")
        return [
            FetchedPayload(
                content=response.content,
                url=str(response.url),
                fetched_at=datetime.now(UTC),
                original_format="xlsx",
                response_headers=self.selected_headers(response),
                published_at=self.header_datetime(response, "last-modified"),
                persist_raw=True,
                resource_version="2026-09",
            )
        ]

    def parse(self, payload: FetchedPayload) -> list[InfrastructureRecord]:
        workbook = load_workbook(io.BytesIO(payload.read_bytes()), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        next(rows, None)
        headers = [normalize_name(str(value or "")) for value in next(rows)]
        if headers[:3] != ["UF", "MUNICIPIO", "TOTAL"] or len(headers) != 24:
            raise SourceValidationError(
                "Cabeçalho SENATRAN alterado",
                code="schema_drift",
                details={"header": headers},
            )
        records: list[InfrastructureRecord] = []
        for values in rows:
            uf, name = (
                str(values[0] or "").strip().upper(),
                normalize_name(str(values[1] or "")),
            )
            municipality = self.municipality_names.get((uf, name))
            if municipality is None:
                continue
            for index, header in enumerate(headers[2:], start=2):
                raw = values[index]
                try:
                    value = None if raw in {None, ""} else Decimal(str(raw))
                except InvalidOperation as exc:
                    raise SourceValidationError(
                        "Quantidade SENATRAN inválida", code="invalid_numeric"
                    ) from exc
                code = (
                    "fleet_total"
                    if header == "TOTAL"
                    else f"fleet_{header.lower().replace(' ', '_')}"
                )
                records.append(
                    InfrastructureRecord(
                        municipality,
                        date(2026, 9, 30),
                        code,
                        "Frota total" if header == "TOTAL" else header.title(),
                        value,
                        payload.url,
                    )
                )
        workbook.close()
        return records

    def validate(self, records: list[InfrastructureRecord]) -> None:
        if not records:
            raise SourceValidationError("Frota SENATRAN vazia", code="empty_dataset")
        grouped: dict[str, dict[str, Decimal]] = defaultdict(dict)
        for record in records:
            if record.value is not None and record.value < 0:
                raise SourceValidationError("Frota negativa", code="out_of_range")
            grouped[record.municipality_ibge_code][record.service_code] = record.value or Decimal(0)
        for municipality, values in grouped.items():
            if values.get("fleet_total") != sum(
                value for code, value in values.items() if code != "fleet_total"
            ):
                raise SourceValidationError(
                    "Total da frota não concilia com tipos",
                    code="reconciliation_failed",
                    details={"municipality": municipality},
                )

    def upsert(
        self, db: Session, records: list[InfrastructureRecord], snapshot: SourceSnapshot
    ) -> None:
        services: dict[str, InfrastructureService] = {}
        for record in records:
            service = services.get(record.service_code) or db.scalar(
                select(InfrastructureService).where(
                    InfrastructureService.code == record.service_code
                )
            )
            if service is None:
                service = InfrastructureService(
                    id=uuid.uuid4(),
                    code=record.service_code,
                    label=record.service_label,
                    service_type="transport",
                )
                db.add(service)
                db.flush()
            services[record.service_code] = service
            key = f"{record.municipality_ibge_code}:2026-09:{record.service_code}"
            previous = db.scalar(
                select(InfrastructureFact)
                .where(
                    InfrastructureFact.source_system == "senatran",
                    InfrastructureFact.source_key == key,
                )
                .order_by(desc(InfrastructureFact.revision))
                .limit(1)
            )
            if previous and previous.value == record.value:
                continue
            db.add(
                InfrastructureFact(
                    id=uuid.uuid4(),
                    source_system="senatran",
                    source_key=key,
                    municipality_ibge_code=record.municipality_ibge_code,
                    reference_date=record.reference_date,
                    service_id=service.id,
                    metric_code="registered_fleet",
                    value=record.value,
                    unit="veículos registrados",
                    availability_status="available" if record.value is not None else "unavailable",
                    source_document_url=record.source_url,
                    snapshot_id=snapshot.id,
                    revision=1 if previous is None else previous.revision + 1,
                )
            )
        db.flush()
        dataset = db.get(Dataset, snapshot.dataset_id)
        if dataset is not None:
            dataset.publication_status = "published"
