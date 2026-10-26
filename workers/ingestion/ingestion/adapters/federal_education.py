from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import unicodedata
import uuid
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, ClassVar
from urllib.parse import urlparse

import httpx
import rarfile
from openpyxl import load_workbook
from pyxlsb import open_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from tenacity import (
    retry,
    retry_if_exception_type,
    stop_after_attempt,
    wait_exponential,
)

from ingestion.adapters.base import USER_AGENT, FetchedPayload, SourceValidationError
from ingestion.adapters.education import MetricRecord, _InepDownloadAdapter
from lente_api.models import EducationMetric, SourceSnapshot

SAEB_2023_URL = "https://download.inep.gov.br/saeb/resultados/planilha_de_resultados_2023.rar"
PNAE_2024_URL = (
    "https://www.gov.br/fnde/pt-br/acesso-a-informacao/acoes-e-programas/programas/"
    "pnae/consultas/repasses-financeiros-por-entidade-executora/"
    "DadosFinanceirosdoPNAE_RedesEstadual_Distrital_Municipal_"
    "PorEntidadeExecutora_2024.xlsx/@@download/file"
)
MISSING = {"", "-", "...", "NA", "N/A", "NULL"}


def normalize_fnde_municipality_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"^PREF(?:EITURA)?\s+MUN(?:ICIPAL)?\s+DE\s+", "", normalized.upper())
    return re.sub(r"[^A-Z0-9]+", " ", normalized).strip()


def _decimal(value: Any, *, brazilian: bool = False) -> Decimal | None:
    if value is None or str(value).strip().upper() in MISSING:
        return None
    raw = str(value).strip()
    if brazilian:
        raw = raw.replace(".", "").replace(",", ".")
    else:
        raw = raw.replace(",", ".")
    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise SourceValidationError(
            f"Valor numérico inválido: {value!r}", code="education_invalid_decimal"
        ) from exc


class _FederalEducationDownloadAdapter(_InepDownloadAdapter):
    allowed_hosts: ClassVar[set[str]]
    download_suffix: str
    original_format: str

    @retry(
        retry=retry_if_exception_type((httpx.TimeoutException, httpx.TransportError)),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        reraise=True,
    )
    def discover(self) -> dict[str, Any]:
        self._validate_url(self.source_url)
        with (
            httpx.Client(
                headers={"User-Agent": USER_AGENT},
                follow_redirects=True,
                timeout=60,
                verify=self._tls_context(),
            ) as client,
            client.stream("GET", self.source_url, headers={"Range": "bytes=0-0"}) as response,
        ):
            response.raise_for_status()
            self._validate_url(str(response.url), redirect=True)
            content_range = response.headers.get("content-range", "")
            total_from_range = content_range.rsplit("/", 1)[-1] if "/" in content_range else ""
            length = int(
                total_from_range
                if total_from_range.isdigit()
                else response.headers.get("content-length", "0")
            )
            if length > self.max_download_bytes:
                raise SourceValidationError(
                    "Arquivo oficial excede o limite", code="download_too_large"
                )
            return {
                "url": str(response.url),
                "status": response.status_code,
                "content_type": response.headers.get("content-type"),
                "content_length": length,
            }

    @retry(
        retry=retry_if_exception_type((httpx.TimeoutException, httpx.TransportError)),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=1, max=8),
        reraise=True,
    )
    def fetch(self) -> list[FetchedPayload]:
        stage_dir = self.raw_data_dir / "_stage"
        stage_dir.mkdir(parents=True, exist_ok=True)
        file_handle, temporary_name = tempfile.mkstemp(
            prefix=f"{self.dataset_code}-", suffix=self.download_suffix, dir=stage_dir
        )
        os.close(file_handle)
        target = Path(temporary_name)
        fetched_at = datetime.now(UTC)
        headers: dict[str, str] = {}
        total = 0
        try:
            with (
                httpx.Client(
                    headers={"User-Agent": USER_AGENT},
                    follow_redirects=True,
                    timeout=httpx.Timeout(300),
                    verify=self._tls_context(),
                ) as client,
                client.stream("GET", self.source_url) as response,
            ):
                response.raise_for_status()
                self._validate_url(str(response.url), redirect=True)
                headers = self.selected_headers(response)
                with target.open("wb") as stream:
                    for chunk in response.iter_bytes(1024 * 1024):
                        total += len(chunk)
                        if total > self.max_download_bytes:
                            raise SourceValidationError(
                                "Arquivo oficial excede o limite",
                                code="download_too_large",
                            )
                        stream.write(chunk)
            self._validate_download(target)
            return [
                FetchedPayload(
                    content=None,
                    file_path=target,
                    url=self.source_url,
                    fetched_at=fetched_at,
                    original_format=self.original_format,
                    response_headers=headers,
                    published_at=self.header_datetime_from_headers(headers, "last-modified"),
                )
            ]
        except Exception:
            target.unlink(missing_ok=True)
            raise

    def _validate_url(self, url: str, *, redirect: bool = False) -> None:
        parsed = urlparse(url)
        if parsed.scheme != "https" or parsed.hostname not in self.allowed_hosts:
            code = "unsafe_redirect" if redirect else "unsafe_source_url"
            raise SourceValidationError("Host do recurso oficial não permitido", code=code)

    def _validate_download(self, path: Path) -> None:
        raise NotImplementedError

    def _upsert_metrics(
        self,
        db: Session,
        records: list[MetricRecord],
        snapshot: SourceSnapshot,
        year: int,
    ) -> None:
        release = self._release(db, snapshot, year)
        if db.scalar(
            select(EducationMetric.id).where(EducationMetric.release_id == release.id).limit(1)
        ):
            return
        for record in records:
            db.add(
                EducationMetric(
                    id=uuid.uuid4(),
                    municipality_ibge_code=record.municipality_code,
                    reference_year=record.year,
                    metric_code=record.metric_code,
                    network_code=record.network_code,
                    stage_code=record.stage_code,
                    value=record.value,
                    availability_status=record.availability_status,
                    snapshot_id=snapshot.id,
                    release_id=release.id,
                )
            )
        db.flush()


class InepSaebAdapter(_FederalEducationDownloadAdapter):
    source_name = "Inep — Saeb municipal 2023"
    dataset_code = "inep_saeb_municipal"
    source_url = SAEB_2023_URL
    year = 2023
    allowed_hosts: ClassVar[set[str]] = {"download.inep.gov.br"}
    download_suffix = ".rar"
    original_format = "rar/xlsb"
    max_download_bytes = 100 * 1024 * 1024
    networks: ClassVar[dict[str, str]] = {
        "Total - Federal, Estadual, Municipal e Privada": "all",
        "Federal": "federal",
        "Estadual": "state",
        "Municipal": "municipal",
    }
    metrics: ClassVar[dict[str, tuple[str, str]]] = {
        "MEDIA_5_LP": ("saeb_portuguese", "fundamental_year_5"),
        "MEDIA_5_MT": ("saeb_mathematics", "fundamental_year_5"),
        "MEDIA_9_LP": ("saeb_portuguese", "fundamental_year_9"),
        "MEDIA_9_MT": ("saeb_mathematics", "fundamental_year_9"),
        "MEDIA_12_LP": ("saeb_portuguese", "secondary_year_3"),
        "MEDIA_12_MT": ("saeb_mathematics", "secondary_year_3"),
    }

    def _validate_download(self, path: Path) -> None:
        with path.open("rb") as stream:
            magic = stream.read(8)
        if not (magic.startswith(b"Rar!\x1a\x07\x00") or magic == b"Rar!\x1a\x07\x01\x00"):
            raise SourceValidationError("Recurso do Saeb não é RAR", code="unexpected_format")
        with rarfile.RarFile(path) as archive:
            members = archive.infolist()
            if len(members) > 20 or sum(item.file_size for item in members) > 1024**3:
                raise SourceValidationError("RAR excede limites seguros", code="unsafe_archive")
            if any(
                Path(item.filename).is_absolute() or ".." in Path(item.filename).parts
                for item in members
            ):
                raise SourceValidationError("Caminho inseguro no RAR", code="unsafe_archive")
            if len([item for item in members if item.filename.lower().endswith(".xlsb")]) != 1:
                raise SourceValidationError(
                    "Planilha municipal do Saeb não identificada",
                    code="saeb_schema_changed",
                )

    @contextmanager
    def _workbook_path(self, payload: FetchedPayload) -> Iterator[Path]:
        archive_path = payload.file_path
        temporary_archive: Path | None = None
        if archive_path is None:
            handle, name = tempfile.mkstemp(suffix=".rar")
            os.close(handle)
            temporary_archive = Path(name)
            temporary_archive.write_bytes(payload.read_bytes())
            archive_path = temporary_archive
        self._validate_download(archive_path)
        extractor = shutil.which("bsdtar") or shutil.which("tar")
        if extractor is None:
            raise SourceValidationError(
                "Extrator libarchive indisponível", code="saeb_extractor_unavailable"
            )
        try:
            with tempfile.TemporaryDirectory(prefix="lente-saeb-") as directory:
                result = subprocess.run(
                    [extractor, "-xf", str(archive_path), "-C", directory],
                    capture_output=True,
                    check=False,
                    timeout=120,
                )
                if result.returncode:
                    raise SourceValidationError(
                        "Falha ao extrair o RAR oficial do Saeb",
                        code="saeb_extract_failed",
                    )
                root = Path(directory).resolve()
                files = [item for item in root.rglob("*") if item.is_file()]
                if any(root not in item.resolve().parents for item in files):
                    raise SourceValidationError("Extração saiu do staging", code="unsafe_archive")
                workbooks = [item for item in files if item.suffix.lower() == ".xlsb"]
                if len(workbooks) != 1:
                    raise SourceValidationError(
                        "Planilha municipal do Saeb não identificada",
                        code="saeb_schema_changed",
                    )
                yield workbooks[0]
        finally:
            if temporary_archive is not None:
                temporary_archive.unlink(missing_ok=True)

    def parse_rows(self, rows: Iterable[list[Any]]) -> list[MetricRecord]:
        iterator = iter(rows)
        headers = [str(value or "").strip() for value in next(iterator, [])]
        required = {
            "ANO_SAEB",
            "CO_MUNICIPIO",
            "DEPENDENCIA_ADM",
            "LOCALIZACAO",
            *self.metrics,
        }
        if not required.issubset(headers):
            raise SourceValidationError("Schema do Saeb alterado", code="saeb_schema_changed")
        records: list[MetricRecord] = []
        for values in iterator:
            row = dict(zip(headers, values, strict=False))
            code = str(row.get("CO_MUNICIPIO") or "").strip().removesuffix(".0")
            network = self.networks.get(str(row.get("DEPENDENCIA_ADM") or "").strip())
            if (
                code not in self.selected_municipalities
                or network is None
                or str(row.get("LOCALIZACAO") or "").strip() != "Total"
            ):
                continue
            year = int(float(row.get("ANO_SAEB") or 0))
            for field, (metric_code, stage_code) in self.metrics.items():
                value = _decimal(row.get(field))
                records.append(
                    MetricRecord(
                        code,
                        year,
                        metric_code,
                        network,
                        stage_code,
                        value,
                        "available" if value is not None else "suppressed",
                    )
                )
        return records

    def parse(self, payload: FetchedPayload) -> list[MetricRecord]:
        with (
            self._workbook_path(payload) as workbook_path,
            open_workbook(str(workbook_path)) as workbook,
        ):
            if len(workbook.sheets) < 4:
                raise SourceValidationError(
                    "Aba municipal do Saeb ausente", code="saeb_schema_changed"
                )
            with workbook.get_sheet(4) as sheet:
                return self.parse_rows([cell.v for cell in row] for row in sheet.rows())

    def validate(self, records: list[MetricRecord]) -> None:
        if not records:
            raise SourceValidationError("Saeb sem linhas para o recorte", code="empty_selection")
        keys = {
            (
                item.municipality_code,
                item.network_code,
                item.stage_code,
                item.metric_code,
            )
            for item in records
        }
        if len(keys) != len(records):
            raise SourceValidationError("Métrica duplicada no Saeb", code="saeb_duplicate_metric")
        if any(
            item.year != self.year or item.value is not None and item.value < 0 for item in records
        ):
            raise SourceValidationError("Valor ou ano inválido no Saeb", code="saeb_invalid_value")

    def upsert(self, db: Session, records: list[MetricRecord], snapshot: SourceSnapshot) -> None:
        self._upsert_metrics(db, records, snapshot, self.year)


class FndePnaeAdapter(_FederalEducationDownloadAdapter):
    source_name = "FNDE — previsão de repasses PNAE 2024"
    dataset_code = "fnde_pnae_municipal"
    source_url = PNAE_2024_URL
    year = 2024
    allowed_hosts: ClassVar[set[str]] = {"www.gov.br"}
    download_suffix = ".xlsx"
    original_format = "xlsx"
    max_download_bytes = 25 * 1024 * 1024
    stages: ClassVar[dict[str, str]] = {
        "PNACN": "pnae_daycare",
        "PNAPN": "pnae_pre_school",
        "PNAFN": "pnae_fundamental",
        "PNAMN": "pnae_secondary",
        "PNAEE": "pnae_special_education",
        "PNAEN": "pnae_youth_adult",
        "PNAIN": "pnae_indigenous",
        "PNAQN": "pnae_quilombola",
        "PNAMI": "pnae_full_time_secondary",
    }

    def __init__(
        self,
        *,
        raw_data_dir: Path,
        selected_municipalities: set[str],
        municipality_names: dict[tuple[str, str], str],
    ):
        super().__init__(raw_data_dir=raw_data_dir, selected_municipalities=selected_municipalities)
        self.municipality_names = municipality_names

    def _validate_download(self, path: Path) -> None:
        with path.open("rb") as stream:
            magic = stream.read(2)
        if magic != b"PK":
            raise SourceValidationError("Recurso do PNAE não é XLSX", code="unexpected_format")
        self._validate_zip(path)

    def parse(self, payload: FetchedPayload) -> list[MetricRecord]:
        source: Path | io.BytesIO = payload.file_path or io.BytesIO(payload.read_bytes())
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers: list[str] | None = None
            records: list[MetricRecord] = []
            for index, values in enumerate(rows, start=1):
                if index == 5:
                    headers = [str(value or "").strip() for value in values[:5]]
                    if headers != [
                        "UF",
                        "ENTIDADE_RESP",
                        "ACAO",
                        "QTDE_ALUNO",
                        "VALOR_PREVISTO",
                    ]:
                        raise SourceValidationError(
                            "Schema do PNAE alterado", code="pnae_schema_changed"
                        )
                    continue
                if index <= 5 or headers is None:
                    continue
                row = dict(zip(headers, values[:5], strict=False))
                state = str(row.get("UF") or "").strip().upper()
                entity = str(row.get("ENTIDADE_RESP") or "").strip()
                municipality = self.municipality_names.get(
                    (state, normalize_fnde_municipality_name(entity))
                )
                action = str(row.get("ACAO") or "").strip().upper()
                stage = self.stages.get(action)
                if municipality not in self.selected_municipalities or stage is None:
                    continue
                students = _decimal(row.get("QTDE_ALUNO"), brazilian=True)
                planned = _decimal(row.get("VALOR_PREVISTO"), brazilian=True)
                for metric_code, value in (
                    ("pnae_students_planned", students),
                    ("pnae_transfer_planned_brl", planned),
                ):
                    records.append(
                        MetricRecord(
                            municipality,
                            self.year,
                            metric_code,
                            "municipal",
                            stage,
                            value,
                            "available" if value is not None else "unavailable",
                        )
                    )
            return records
        finally:
            workbook.close()

    def validate(self, records: list[MetricRecord]) -> None:
        if not records:
            raise SourceValidationError("PNAE sem linhas para o recorte", code="empty_selection")
        keys = {
            (
                item.municipality_code,
                item.metric_code,
                item.network_code,
                item.stage_code,
            )
            for item in records
        }
        if len(keys) != len(records):
            raise SourceValidationError("Métrica duplicada no PNAE", code="pnae_duplicate_metric")
        if any(
            item.year != self.year or item.value is not None and item.value < 0 for item in records
        ):
            raise SourceValidationError("Valor ou ano inválido no PNAE", code="pnae_invalid_value")

    def upsert(self, db: Session, records: list[MetricRecord], snapshot: SourceSnapshot) -> None:
        self._upsert_metrics(db, records, snapshot, self.year)
